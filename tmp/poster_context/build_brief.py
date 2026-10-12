from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
out = root / 'output' / 'documents' / 'AutoMetric_AI_Poster_Context.docx'
out.parent.mkdir(parents=True, exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'List Bullet']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    for b in list(st.element.iter(qn('w:pBdr'))):
        b.getparent().remove(b)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(5)
d.styles['Normal'].paragraph_format.line_spacing = 1.03
d.styles['Title'].font.size = Pt(23)
d.styles['Title'].paragraph_format.space_after = Pt(6)
d.styles['Subtitle'].font.size = Pt(12)
d.styles['Heading 1'].font.size = Pt(13)
d.styles['Heading 1'].paragraph_format.space_before = Pt(9)
d.styles['Heading 1'].paragraph_format.space_after = Pt(4)

def h(t): d.add_heading(t, 1)
def p(t): d.add_paragraph(t)
def b(t): d.add_paragraph(t, 'List Bullet')
def step(label, text):
    x = d.add_paragraph()
    x.add_run(label + ' ').bold = True
    x.add_run(text)

d.add_paragraph('AutoMetric AI Poster Context', 'Title')
d.add_paragraph('Used car price estimation and maintenance management in Pakistan', 'Subtitle')
p('A concise content bank connecting data collection, machine learning, Flutter, Firebase, FastAPI, and Docker. The sections match the supplied FYP poster structure and can be shortened when assembling the poster.')
h('Abstract')
p('AutoMetric AI combines machine learning based used car price estimation with vehicle maintenance management for Pakistan. PakWheels listings were collected and processed into a structured dataset. Regression models were compared, and an XGBoost model using 25 engineered features was integrated with a FastAPI backend. A responsive Flutter application supports Android and web, displays estimates in PKR, and uses Firebase for accounts and personal records. Docker packages the prediction service and model artifacts for consistent execution.')
h('Problem Statement')
p('Used car asking prices vary with brand, variant, age, mileage, and specifications, making comparable valuations difficult. Vehicle owners also need an organized way to record services and track upcoming maintenance.')
h('Aim and Objectives')
p('Develop an Android and web application that estimates Pakistani used car prices and helps owners manage vehicle maintenance.')
b('Collect and clean local listing data, engineer vehicle features, and compare regression models.')
b('Connect a price prediction API to an accessible Flutter interface.')
b('Support saved valuations, separate vehicle records, service expenses, and maintenance reminders.')
h('Proposed Solution')
p('Guests can estimate prices without an account. Signed-in users can save evaluations and manage a separate garage with multiple cars, odometer updates, service records, PKR costs, and date or kilometre based schedules. Maintenance records are not linked to prediction history.')
h('Methodology')
step('1 Collect', 'Gather PakWheels listing URLs and vehicle details using Requests and BeautifulSoup. Merge CSV batches and remove duplicate listing IDs. A separate scraper also collects images using Selenium; images are not model inputs.')
step('2 Prepare', 'Standardize names and units, handle missing values, and retain valid price records. Expand 18 features to 25 using hierarchical brand-model-trim encoding, age, mileage rates, variant flags, and interactions.')
step('3 Train', 'Compare Ridge, Random Forest, XGBoost, LightGBM, and CatBoost. Use Optuna in tuning experiments; train the production XGBoost v3 with a revised fixed configuration, log-transformed prices, and an 80/20 split.')
step('4 Integrate', 'Export model and preprocessing artifacts, load them once at FastAPI startup, and connect Flutter through JSON requests. Firebase stores user records; Docker packages the backend.')

d.add_page_break()
d.add_paragraph('Results and Technical Context', 'Title')
h('Results')
b('Dataset: 10,008 raw rows; 9,849 processed records; 7,879 training and 1,970 test records.')
b('Saved XGBoost v3 metrics: R squared 0.8646; MAE 4.84 lakhs PKR; MAPE 10.06%; RMSE 22.52 lakhs PKR.')
b('XGBoost v3 model package is approximately 2.94 MB, compared with approximately 73.33 MB for the tuned v3 Random Forest package.')
p('Evaluation qualification: Price-based target encodings were calculated before the train/test split. These scores are therefore preliminary and require re-evaluation with training-only or out-of-fold encodings. R squared is not percentage accuracy. The dataset represents advertised asking prices, not verified sale prices.')
h('Technologies Used')
step('Data', 'Python, Requests, BeautifulSoup, Selenium, Pandas, and CSV.')
step('Machine learning', 'NumPy, scikit-learn, XGBoost, and Joblib; Optuna, LightGBM, and CatBoost for experimentation.')
step('Application', 'Flutter, Dart, Material UI, Firebase Authentication, and Cloud Firestore on the Spark plan.')
step('Backend', 'FastAPI, Pydantic, Uvicorn, and Docker using a Python 3.11 slim image.')
h('System Workflow')
p('Training: PakWheels listings → merged CSV → cleaning and feature engineering → model comparison and training → exported model and preprocessing artifacts.')
p('Prediction: Flutter form with 11 vehicle inputs → FastAPI validation and 25-feature preparation → XGBoost → estimated price in lakhs PKR → result and optional saved evaluation.')
p('Maintenance: Separate garage → service and odometer records in Firestore → date and distance due calculations → reminder centre and platform alerts.')
h('Maintenance and Deployment Scope')
p('Service intervals are user-defined. Confirmed odometer readings determine distance-based due status; driving-rate forecasts are estimates, not live tracking. Android uses local scheduled notifications. Web provides in-app and optional foreground browser alerts, not reminders after the site closes. This avoids a paid server notification service on Spark.')
p('The current Flutter configuration uses a local API address. Docker packaging provides reproducible execution but does not establish public cloud deployment. Production rollout includes security-rule deployment and real-device notification checks.')
h('Summary and Project Demo')
p('AutoMetric AI connects a local listing dataset and an engineered predictor to a cross-platform app, alongside personal vehicle maintenance management. Demonstrate a price prediction, saved history, adding a car and service, an odometer update, and the reminder centre. Add the actual demo QR code, team names, and supervisor details when composing the poster.')
h('Source Map')
p('Data: D:\\Data Scrapping\\scrapping and excels. Models: D:\\Model Training\\pipeline_v3_production and selected_models_experiment. Backend: D:\\AutoMetricAI-v3-Backend\\app and Dockerfile. App: D:\\Flutter Projects\\autometric_ai\\lib and docs\\maintenance.md.')
d.core_properties.title = 'AutoMetric AI Poster Context'
d.core_properties.subject = 'Concise FYP project context and poster content'
d.core_properties.author = 'AutoMetric AI Project'
d.save(out)
print(out)
print('Words:', sum(len(x.text.split()) for x in d.paragraphs))
